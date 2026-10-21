import tempfile
import unittest
from pathlib import Path

from docx import Document
from clean_docx_batch import extract_clean_script
from clean_library_audio import LIBRARY_END, run, source_matches


class LibraryCleaningTests(unittest.TestCase):
    def test_workshop_heading_cuts_off_exercises_but_not_teaching_mentions(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'lesson.docx'
            document = Document()
            document.add_paragraph('L.L INVEST IN YOURSELF 3 PAGES')
            document.add_paragraph('Volunteering and attending workshops help you learn.')
            document.add_paragraph('Keep investing in yourself.')
            document.add_paragraph('Workshop Objective:\nComplete these group exercises.')
            document.add_paragraph('My to do list: Do not narrate this.')
            document.save(source)
            text, cutoff = extract_clean_script(source, end_marker=LIBRARY_END)
            self.assertTrue(cutoff)
            self.assertEqual(text, 'Volunteering and attending workshops help you learn.\n\nKeep investing in yourself.')

    def test_explicit_aliases_preserve_future_versions_and_avoid_unrelated_work_lesson(self):
        future1, future2 = Path('future1.docx'), Path('future2.docx')
        work, acronym = Path('L.L WORK 2 PAGES.docx'), Path('L.L W.O.R.K 2 PAGES.docx')
        lookup = {'futurev1': [future1], 'futurev2': [future2], 'work': [work, acronym]}
        self.assertEqual(source_matches('Future', lookup), [future1, future2])
        self.assertEqual(source_matches('W.O.R.K. - The Perspective You NEVER Had on WORK', lookup), [acronym])
        self.assertEqual(source_matches('A completely new unmapped session', lookup), [])

    def test_existing_output_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            sentinel = output / 'existing.txt'
            sentinel.write_text('Keep this file.')
            with self.assertRaisesRegex(ValueError, 'already exists'):
                run(output / 'input', output / 'contents.docx', output)
            self.assertEqual(sentinel.read_text(), 'Keep this file.')


if __name__ == '__main__':
    unittest.main()
