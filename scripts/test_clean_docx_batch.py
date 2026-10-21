"""Run with: python -m unittest discover -s scripts -p test_clean_docx_batch.py"""
import contextlib
import hashlib
import io
from pathlib import Path
import tempfile
import unittest

from docx import Document
from clean_docx_batch import extract_clean_script, main


class CleanDocxTests(unittest.TestCase):
    def test_notes_and_goals_cutoffs(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'lesson.docx'
            for marker in ['Notes:', 'My Goals today are:', 'My to do list:']:
                document = Document()
                document.add_paragraph('Read this.')
                document.add_paragraph('-----_____')
                document.add_paragraph(marker)
                document.add_paragraph('Do not include this worksheet.')
                document.save(source)
                self.assertEqual(extract_clean_script(source), ('Read this.', True))

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def test_headers_cues_and_inline_cutoff_preserve_spoken_body(self):
        source = self.root / '2.1-IDENTITY.docx'
        document = Document()
        document.sections[0].header.paragraphs[0].text = 'Private document header'
        document.sections[0].footer.paragraphs[0].text = 'Page footer'
        document.add_paragraph('SECTION 1: IDENTITY', style='Heading 3')
        document.add_paragraph('[Tone: Warm]')
        document.add_paragraph('IDENTITY.')
        document.add_paragraph('Welcome.')
        document.add_paragraph('WE ALL HAVE UNIQUE GIFTS.')
        document.add_paragraph('A closing sentence. My TO-DO LIST: Questions to complete.')
        document.add_paragraph('Never narrate this workshop.')
        document.save(source)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        text, found_end = extract_clean_script(source)
        self.assertTrue(found_end)
        self.assertEqual(text, 'Welcome.\n\nWE ALL HAVE UNIQUE GIFTS.\n\nA closing sentence.')
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_tables_and_multiline_markers_preserve_document_order(self):
        source = self.root / 'hope.docx'
        document = Document()
        document.add_paragraph('L.L HOPE 2 PAGES')
        document.add_paragraph("Let's talk about hope.")
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'A spoken line in a table.'
        table.cell(1, 0).text = 'Last spoken line.\nMy to do list this week:\nDo not read.'
        document.add_paragraph('Also do not read.')
        document.save(source)
        text, found_end = extract_clean_script(source)
        self.assertTrue(found_end)
        self.assertEqual(text, "Let's talk about hope.\n\nA spoken line in a table.\n\nLast spoken line.")

    def test_no_marker_keeps_the_whole_body(self):
        source = self.root / 'lesson.docx'
        document = Document()
        document.add_paragraph('Welcome.')
        document.add_paragraph('A sentence about things to do with your gifts.')
        document.add_paragraph('Keep this ending.')
        document.save(source)
        text, found_end = extract_clean_script(source)
        self.assertFalse(found_end)
        self.assertIn('Keep this ending.', text)

    def test_recursive_batch_keeps_duplicate_names_separate_and_skips_word_locks(self):
        for folder in ['module-a', 'module-b']:
            source = self.root / 'input' / folder / 'lesson.docx'
            source.parent.mkdir(parents=True)
            document = Document()
            document.add_paragraph('Shared document subtitle', style='Title' if folder == 'module-a' else None)
            document.add_paragraph(f'Spoken body for {folder}.')
            document.add_paragraph('My to do list: Stop here.')
            document.save(source)
        (self.root / 'input' / '~$lesson.docx').write_bytes(b'Not a real document')
        output = self.root / 'output'
        with contextlib.redirect_stdout(io.StringIO()):
            result = main(['--input', str(self.root / 'input'), '--output', str(output)])
        self.assertEqual(result, 0)
        self.assertEqual(len(list(output.rglob('*.txt'))), 2)
        self.assertEqual((output / 'module-a' / 'lesson_clean.txt').read_text(), 'Spoken body for module-a.\n')
        self.assertEqual((output / 'module-b' / 'lesson_clean.txt').read_text(), 'Spoken body for module-b.\n')


if __name__ == '__main__':
    unittest.main()
